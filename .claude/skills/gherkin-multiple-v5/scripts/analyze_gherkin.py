#!/usr/bin/env python3
import argparse
import json
import re
import sys
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
import zipfile
from typing import Dict, List, Tuple

STEP_KEYWORDS = ("Given", "When", "Then", "And", "But", "*")
HEAVY_BG_HINTS = {
    "create", "insert", "seed", "load", "import", "call", "post", "put", "delete", "upload", "generate"
}
QUOTE_RE = re.compile(r'"[^"]*"|\'[^\']*\'')
NUMBER_RE = re.compile(r"\b\d+\b")
PLACEHOLDER_RE = re.compile(r"<[^>]+>")
TITLE_RE = re.compile(r"^(Feature|Rule|Scenario Outline|Scenario|Background|Examples):\s*(.*)$")
VALID_SUFFIXES = {".feature", ".md", ".txt"}
STATE_VERSION = 9

# LKS Next working palette for Excel outputs.
LKS_NAVY = "003B5C"
LKS_CYAN = "00A3E0"
LKS_CHARCOAL = "2E3135"
LKS_LIGHT_GRAY = "F3F5F7"
LKS_MID_GRAY = "B7C1CA"
LKS_WHITE = "FFFFFF"
LKS_GREEN = "84BD00"
LKS_ORANGE = "F28C28"



def strip_comments(line: str) -> str:
    if "#" in line:
        return line.split("#", 1)[0].rstrip()
    return line.rstrip()


def normalize_step_text(text: str) -> str:
    t = QUOTE_RE.sub('"{string}"', text)
    t = PLACEHOLDER_RE.sub('{param}', t)
    t = NUMBER_RE.sub('{int}', t)
    return re.sub(r"\s+", " ", t).strip().lower()


def normalize_feature_key(text: str) -> str:
    base = re.sub(r"\s+", " ", (text or "").strip().lower())
    return base or "untitled-feature"


def strip_report_for_state(report: Dict) -> Dict:
    clean = json.loads(json.dumps(report))
    clean.pop("state_feature_key", None)
    return clean


def clamp(value: float, lo: float = 0.0, hi: float = 100.0) -> float:
    return max(lo, min(hi, value))


def score_positive(value: float, target: float) -> float:
    return 100.0 if target <= 0 else clamp(100.0 * value / target)


def score_negative(value: float, bad_threshold: float) -> float:
    return 100.0 if bad_threshold <= 0 else clamp(100.0 * (1.0 - value / bad_threshold))


def score_asl(asl: float) -> float:
    if asl <= 5:
        return 100.0
    if asl >= 7:
        return 0.0
    return 100.0 * (7 - asl) / 2.0


def safe_div(numerator: float, denominator: float, default: float) -> float:
    return default if denominator == 0 else numerator / denominator


def fmt_num(value, decimals: int = 4) -> str:
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        text = f"{value:.{decimals}f}"
        return text.rstrip('0').rstrip('.') if '.' in text else text
    return str(value)


def table_safe(text) -> str:
    return str(text).replace('|', '\\|').replace('\n', '<br>')


def metric_row(metric, scope, what, how, formula, calculated_data, raw_value, normalized_score, example, raw_data=''):
    return {
        'metric': metric,
        'concept_scope': scope,
        'what_it_measures': what,
        'how_to_measure_it': how,
        'practical_formula_rule': formula,
        'calculated_data': calculated_data,
        'raw_value': round(raw_value, 4) if isinstance(raw_value, float) else raw_value,
        'normalized_score': round(normalized_score, 1) if isinstance(normalized_score, float) else normalized_score,
        'example': example,
        'raw_data': raw_data,
    }


def interpretation_label(score: float) -> str:
    return 'strong' if score >= 75 else 'moderate' if score >= 50 else 'weak'


def parse_gherkin(text: str) -> Dict:
    feature = {"title": None, "background_steps": [], "background_count": 0, "scenarios": []}
    current = None
    in_examples = False
    example_headers: List[str] = []
    example_rows: List[Dict[str, str]] = []

    for raw in text.splitlines():
        line = strip_comments(raw).strip()
        if not line:
            continue
        match = TITLE_RE.match(line)
        if match:
            kind, title = match.group(1), match.group(2).strip()
            if kind == 'Feature':
                feature['title'] = title
                current = None
                in_examples = False
            elif kind == 'Background':
                feature['background_count'] += 1
                current = {'kind': 'Background', 'title': title or 'Background', 'steps': []}
                in_examples = False
            elif kind in ('Scenario', 'Scenario Outline'):
                if current and current.get('kind') in ('Scenario', 'Scenario Outline'):
                    current['example_headers'] = example_headers
                    current['example_rows'] = example_rows
                    feature['scenarios'].append(current)
                current = {'kind': kind, 'title': title, 'steps': []}
                in_examples = False
                example_headers = []
                example_rows = []
            elif kind == 'Examples':
                in_examples = True
            continue

        if line.startswith('|') and in_examples and current and current.get('kind') == 'Scenario Outline':
            cells = [cell.strip() for cell in line.strip('|').split('|')]
            if not example_headers:
                example_headers = cells
            elif len(cells) == len(example_headers):
                example_rows.append(dict(zip(example_headers, cells)))
            continue

        if any(line.startswith(k + ' ') for k in STEP_KEYWORDS[:-1]) or line.startswith('* '):
            parts = line.split(maxsplit=1)
            keyword = parts[0]
            text_part = parts[1] if len(parts) > 1 else ''
            step = {'keyword': keyword, 'text': text_part, 'normalized': normalize_step_text(text_part)}
            if current and current.get('kind') == 'Background':
                feature['background_steps'].append(step)
            elif current and current.get('kind') in ('Scenario', 'Scenario Outline'):
                current['steps'].append(step)
            continue

    if current and current.get('kind') in ('Scenario', 'Scenario Outline'):
        current['example_headers'] = example_headers
        current['example_rows'] = example_rows
        feature['scenarios'].append(current)
    return feature


def format_candidate_reusable_raw_data(details: List[Dict]) -> str:
    if not details:
        return 'None. No candidate reusable families detected.'
    parts = []
    for index, detail in enumerate(details, start=1):
        variants = ' || '.join(detail['steps'])
        parts.append(f"{index}. {detail['normalized_family']} => {variants}")
    return '; '.join(parts)


def format_duplicated_non_outline_raw_data(details: List[Dict]) -> str:
    if not details:
        return 'None. No duplicated non-outline scenario families detected.'
    parts = []
    for index, detail in enumerate(details, start=1):
        scenario_list = ' || '.join(detail['scenario_titles'])
        normalized_steps = ' -> '.join(detail['normalized_step_sequence'])
        parts.append(f"{index}. scenarios: {scenario_list}; normalized steps: {normalized_steps}")
    return '; '.join(parts)


def collect_background_candidates(scenarios: List[Dict]) -> Dict:
    scenarios = [s for s in scenarios if s.get('steps')]
    shared_step_counter = Counter()
    blocks = []
    max_depth = 0

    def walk(group: List[Dict], depth: int, prefix: List[str]) -> None:
        nonlocal max_depth
        buckets: Dict[str, List[Dict]] = defaultdict(list)
        for scenario in group:
            if depth < len(scenario['steps']):
                buckets[scenario['steps'][depth]['text']].append(scenario)
        shared_children = []
        for text, sub in buckets.items():
            if len(sub) >= 2:
                shared_children.append((text, sub))
        if prefix and not shared_children:
            blocks.append({
                'scenario_titles': [scenario.get('title') or scenario.get('kind') for scenario in group],
                'scenario_count': len(group),
                'block_length': len(prefix),
                'occurrences': len(prefix) * len(group),
                'lines': list(prefix),
            })
            return
        for text, sub in shared_children:
            shared_step_counter[text] += len(sub)
            max_depth = max(max_depth, depth + 1)
            walk(sub, depth + 1, prefix + [text])

    if len(scenarios) >= 2:
        walk(scenarios, 0, [])

    blocks.sort(key=lambda item: (-item['occurrences'], -item['block_length'], ';'.join(item['lines'])))
    return {
        'background_candidate_distinct_lines': len(shared_step_counter),
        'background_candidate_occurrences': sum(shared_step_counter.values()),
        'background_candidate_lines': sorted(shared_step_counter.keys()),
        'background_candidate_blocks': blocks,
        'max_shared_prefix_depth': max_depth,
        'prefix_occurrences_by_step': dict(shared_step_counter),
    }


def collect_outline_candidates(scenarios: List[Dict]) -> Dict:
    regular_scenarios = [scenario for scenario in scenarios if scenario['kind'] == 'Scenario']
    families: Dict[Tuple[str, ...], List[Dict]] = defaultdict(list)
    for scenario in regular_scenarios:
        key = tuple(step['normalized'] for step in scenario['steps'])
        families[key].append(scenario)

    outline_candidate_lines = set()
    outline_candidate_occurrences = 0
    family_rows = []
    duplicated_non_outline_scenarios = []

    for key, family in families.items():
        if len(family) <= 1:
            continue
        duplicated_non_outline_scenarios.append({
            'scenario_titles': [scenario['title'] or scenario['kind'] for scenario in family],
            'normalized_step_sequence': list(key),
        })
        variable_positions = []
        for idx in range(len(key)):
            literals = sorted({scenario['steps'][idx]['text'] for scenario in family if idx < len(scenario['steps'])})
            if len(literals) > 1:
                variable_positions.append({
                    'position': idx + 1,
                    'normalized': key[idx],
                    'literals': literals,
                })
                outline_candidate_lines.update(literals)
                outline_candidate_occurrences += len(family)
        if variable_positions:
            family_rows.append({
                'scenario_titles': [scenario['title'] or scenario['kind'] for scenario in family],
                'scenario_count': len(family),
                'variable_positions': variable_positions,
            })

    family_rows.sort(key=lambda item: (-item['scenario_count'], -len(item['variable_positions'])))
    return {
        'outline_candidate_distinct_lines': len(outline_candidate_lines),
        'outline_candidate_occurrences': outline_candidate_occurrences,
        'outline_candidate_lines': sorted(outline_candidate_lines),
        'outline_candidate_families': family_rows,
        'duplicated_non_outline_scenarios': duplicated_non_outline_scenarios,
    }


def summarize_feature(feature: Dict, source_name: str = '') -> Dict:
    scenarios = feature['scenarios']
    background_steps = feature['background_steps']
    background_count = feature.get('background_count', 0) or (1 if background_steps else 0)
    avg_background_step_count = safe_div(len(background_steps), background_count, 0.0)
    all_steps = [step for scenario in scenarios for step in scenario['steps']]
    outlines = [scenario for scenario in scenarios if scenario['kind'] == 'Scenario Outline']

    step_families: Dict[str, List[str]] = defaultdict(list)
    for step in all_steps:
        step_families[step['normalized']].append(step['text'])

    candidate_reusable = 0
    parameterized_families = 0
    candidate_reusable_details = []
    for normalized_family, originals in step_families.items():
        unique_originals = sorted(set(originals))
        if len(unique_originals) > 1:
            candidate_reusable += 1
            candidate_reusable_details.append({
                'normalized_family': normalized_family,
                'steps': unique_originals,
            })
            if ('{string}' in normalized_family) or ('{int}' in normalized_family) or ('{param}' in normalized_family):
                parameterized_families += 1

    families: Dict[Tuple[str, ...], List[Dict]] = defaultdict(list)
    for scenario in scenarios:
        key = tuple(step['normalized'] for step in scenario['steps'])
        families[key].append(scenario)
    duplicated_families = [family for family in families.values() if len(family) > 1]
    outline_modeled_families = sum(1 for family in duplicated_families if any(s['kind'] == 'Scenario Outline' for s in family))

    total_example_rows = 0
    max_outline_rows = 0
    outline_breakdown = []
    for scenario in outlines:
        rows = len(scenario.get('example_rows', []))
        cols = len(scenario.get('example_headers', []))
        total_example_rows += rows
        max_outline_rows = max(max_outline_rows, rows)
        outline_breakdown.append({
            'title': scenario['title'],
            'rows': rows,
            'columns': cols,
            'product': rows * max(cols, 1),
        })

    background_heavy = any(any(hint in step['text'].lower() for hint in HEAVY_BG_HINTS) for step in background_steps)
    background_overused = bool(background_count and (avg_background_step_count > 4 or background_heavy))
    scenario_lengths = [len(scenario['steps']) for scenario in scenarios]

    step_occurrences = []
    feature_label = feature.get('title') or source_name
    for step in background_steps:
        step_occurrences.append({
            'feature': feature_label,
            'scenario': 'Background',
            'kind': 'Background',
            'step': step['text'],
            'normalized': step['normalized'],
        })

    total_executed_steps = len(background_steps)
    for scenario in scenarios:
        scenario_label = scenario['title'] or scenario['kind']
        total_executed_steps += len(scenario['steps'])
        for step in scenario['steps']:
            step_occurrences.append({
                'feature': feature_label,
                'scenario': scenario_label,
                'kind': scenario['kind'],
                'step': step['text'],
                'normalized': step['normalized'],
            })

    background_candidates = collect_background_candidates(scenarios)
    outline_candidates = collect_outline_candidates(scenarios)
    total_scenario_steps = sum(scenario_lengths)
    mbo = safe_div(background_candidates['background_candidate_occurrences'], total_scenario_steps, 0.0)
    bdi = background_candidates['background_candidate_distinct_lines'] * mbo
    bdi_s = score_negative(bdi, 6.0)

    return {
        'source_name': source_name,
        'feature_title': feature.get('title'),
        'scenario_count': len(scenarios),
        'outline_count': len(outlines),
        'background_count': background_count,
        'background_step_count': len(background_steps),
        'avg_background_step_count': avg_background_step_count,
        'max_background_step_count': avg_background_step_count,
        'total_steps': len(all_steps),
        'total_executed_steps': total_executed_steps,
        'unique_normalized_steps': len(set(step['normalized'] for step in all_steps)),
        'normalized_step_set': sorted(set(step['normalized'] for step in all_steps)),
        'normalized_step_family_count': len(step_families),
        'candidate_reusable_families': candidate_reusable,
        'parameterized_reusable_families': parameterized_families,
        'candidate_reusable_details': candidate_reusable_details,
        'duplicated_families': len(duplicated_families),
        'outline_modeled_duplicated_families': outline_modeled_families,
        'duplicated_non_outline_scenarios': outline_candidates['duplicated_non_outline_scenarios'],
        'total_example_rows': total_example_rows,
        'outline_breakdown': outline_breakdown,
        'background_overused': background_overused,
        'background_overused_features': 1 if background_overused else 0,
        'background_heavy': background_heavy,
        'background_texts': [step['text'] for step in background_steps],
        'total_scenario_steps': total_scenario_steps,
        'scenario_lengths': scenario_lengths,
        'max_outline_rows': max_outline_rows,
        'background_over6': avg_background_step_count > 6,
        'step_occurrences': step_occurrences,
        'background_candidate_distinct_lines': background_candidates['background_candidate_distinct_lines'],
        'background_candidate_occurrences': background_candidates['background_candidate_occurrences'],
        'background_candidate_lines': background_candidates['background_candidate_lines'],
        'background_candidate_blocks': background_candidates['background_candidate_blocks'],
        'max_shared_prefix_depth': background_candidates['max_shared_prefix_depth'],
        'prefix_occurrences_by_step': background_candidates['prefix_occurrences_by_step'],
        'mbo': mbo,
        'bdi': bdi,
        'bdi_s': bdi_s,
    }


def build_feature_report(aggregate: Dict) -> Dict:
    total_scenarios = aggregate['scenario_count']

    spc = safe_div(aggregate['parameterized_reusable_families'], aggregate['candidate_reusable_families'], 1.0)
    spc_s = score_positive(spc, 0.80)
    if aggregate['candidate_reusable_families']:
        spc_calc = (
            f"Normalized step families = {aggregate['normalized_step_family_count']}; candidate reusable families = {aggregate['candidate_reusable_families']}; "
            f"parameterized reusable families = {aggregate['parameterized_reusable_families']}; raw = {aggregate['parameterized_reusable_families']}/{aggregate['candidate_reusable_families']} = {fmt_num(spc)}"
        )
    else:
        spc_calc = (
            f"Normalized step families = {aggregate['normalized_step_family_count']}; candidate reusable families = 0; "
            f"no duplicate literal families were detected, so SPC defaults to 1.0"
        )

    oar = safe_div(aggregate['outline_modeled_duplicated_families'], aggregate['duplicated_families'], 1.0)
    oar_s = score_positive(oar, 0.80)
    if aggregate['duplicated_families']:
        oar_calc = (
            f"Duplicated scenario families = {aggregate['duplicated_families']}; duplicated families modeled as outlines = {aggregate['outline_modeled_duplicated_families']}; "
            f"raw = {aggregate['outline_modeled_duplicated_families']}/{aggregate['duplicated_families']} = {fmt_num(oar)}"
        )
    else:
        oar_calc = 'Duplicated scenario families = 0; no repeated normalized scenario families were found, so OAR defaults to 1.0'

    background_overused = bool(aggregate['background_overused_features'])
    background_overuse_score = 0.0 if background_overused else 100.0
    bg_text = '; '.join(aggregate['background_texts']) if aggregate['background_texts'] else 'none'

    background_debt_raw = aggregate['bdi']
    background_debt_score = aggregate['bdi_s']
    background_design_score = min(background_overuse_score, background_debt_score)
    background_design_raw = f"overused={'yes' if background_overused else 'no'}; debt={fmt_num(background_debt_raw)}"
    background_design_calc = (
        f"Background sections = {aggregate.get('background_count', 0)}; total background steps = {aggregate['background_step_count']}; average background steps = {fmt_num(aggregate.get('avg_background_step_count', 0.0))}; heavy-setup hint detected = {'yes' if aggregate['background_heavy'] else 'no'}; "
        f"background overused = {'yes' if background_overused else 'no'}; overuse_score = {fmt_num(background_overuse_score,1)}; background content: {bg_text}; "
        f"background candidate distinct lines = {aggregate['background_candidate_distinct_lines']}; background candidate occurrences = {aggregate['background_candidate_occurrences']}; total scenario steps = {aggregate['total_scenario_steps']}; "
        f"MBO = {aggregate['background_candidate_occurrences']}/{aggregate['total_scenario_steps'] if aggregate['total_scenario_steps'] else 1} = {fmt_num(aggregate['mbo'])}; background_debt_raw = {aggregate['background_candidate_distinct_lines']} * {fmt_num(aggregate['mbo'])} = {fmt_num(background_debt_raw)}; debt_score = {fmt_num(background_debt_score,1)}; final BDI_s = min(overuse_score, debt_score) = {fmt_num(background_design_score,1)}"
        if aggregate['total_scenario_steps'] or aggregate.get('background_count', 0) else
        'No Background section and no scenario steps; overuse_score = 100; debt_score = 100; final BDI_s = 100'
    )

    asl = safe_div(aggregate['total_scenario_steps'], total_scenarios, 0.0)
    asl_s = score_asl(asl)
    asl_calc = (
        f"Scenario lengths = {aggregate['scenario_lengths']}; total scenario steps = {aggregate['total_scenario_steps']}; total scenarios = {total_scenarios}; raw = {aggregate['total_scenario_steps']}/{total_scenarios if total_scenarios else 1} = {fmt_num(asl)}"
        if total_scenarios else 'Total scenarios = 0; ASL defaults to 0'
    )


    spc_raw_data = format_candidate_reusable_raw_data(aggregate.get('candidate_reusable_details', []))
    oar_raw_data = format_duplicated_non_outline_raw_data(aggregate.get('duplicated_non_outline_scenarios', []))
    background_raw_data = 'None. No repeated initial scenario blocks detected.'
    if aggregate.get('background_candidate_blocks'):
        block_parts = []
        for index, block in enumerate(aggregate['background_candidate_blocks'][:10], start=1):
            lines = ' -> '.join(block['lines'])
            scenarios = ' || '.join(block['scenario_titles'])
            block_parts.append(f"{index}. scenarios: {scenarios}; lines: {lines}")
        background_raw_data = '; '.join(block_parts)

    metrics = [
        metric_row(
            'Parameterization of Steps (SPC)',
            'step',
            'Checks whether repeated step families with data variations are properly parameterized instead of duplicated as separate literal step definitions.',
            'Normalize step text, group equivalent step families, and count how many candidate reusable families are already parameterized.',
            'SPC = parameterized reusable families / candidate reusable families. Target >= 0.80; SPC_s = min(100, 100 * SPC / 0.80).',
            spc_calc,
            spc,
            spc_s,
            'Interpretation: 100 means candidate step families are sufficiently parameterized; low values indicate duplicated steps with embedded literals.',
            spc_raw_data
        ),
        metric_row(
            'Scenario Outline for Data-Driven Rules (OAR)',
            'scenario',
            'Checks whether repeated scenarios that only differ by data are modeled as Scenario Outline with Examples.',
            'Group scenarios by normalized step sequence and verify which duplicated scenario families are represented by Scenario Outline.',
            'OAR = duplicated families modeled as Scenario Outline / total duplicated families. Target >= 0.80; OAR_s = min(100, 100 * OAR / 0.80).',
            oar_calc,
            oar,
            oar_s,
            'Interpretation: 100 means data-driven duplication is well consolidated; low values point to copied scenarios that should probably become Scenario Outlines.',
            oar_raw_data
        ),
        metric_row(
            'Background Design Index (BDI)',
            'feature',
            'Evaluates Background design in one metric: it penalizes both existing Background sections that are too long or heavy and repeated setup that should likely be extracted to Background.',
            'Compute a Background overuse score and a repeated setup debt score. The final score takes the worse signal so one problem cannot hide the other.',
            'overuse_score = 0 if Background is overused, otherwise 100. debt_score = max(0, 100 * (1 - background_debt_raw / 6.0)). BDI_s = min(overuse_score, debt_score).',
            background_design_calc,
            background_design_raw,
            background_design_score,
            'Interpretation: 100 means balanced Background usage and little setup debt; low values indicate heavy Background sections or repeated setup still embedded in scenarios.',
            background_raw_data
        ),
        metric_row(
            'Average Scenario Length (ASL)',
            'feature',
            'Measures average scenario length as a proxy for cognitive load, readability, and maintainability.',
            'Count executable steps in each Scenario and Scenario Outline, excluding Background, and average them across scenarios.',
            'ASL = total scenario steps / total scenarios. Score is 100 when ASL <= 5 and decays linearly to 0 when ASL >= 7.',
            asl_calc,
            asl,
            asl_s,
            'Interpretation: scenarios with 3 to 5 steps are usually easy to maintain; above 7 steps, consider splitting or simplifying them.',
        ),
    ]

    fmi = (spc_s + oar_s + background_design_score + asl_s) / 4.0
    guardrails = []
    if aggregate['background_over6']:
        guardrails.append({'cap': 60.0, 'reason': 'Average Background length is more than 6 steps'})
    final_fmi = min([fmi] + [g['cap'] for g in guardrails]) if guardrails else fmi
    total_row = metric_row(
        'Feature Maintainability Index (FMI)',
        'feature',
        'Final structural maintainability index for the feature, combining step parameterization, Scenario Outline usage, Background design, and scenario length.',
        'Average the normalized SPC_s, OAR_s, BDI_s, and ASL_s scores. Then apply the Background guardrail when applicable.',
        'FMI = (SPC_s + OAR_s + BDI_s + ASL_s) / 4',
        f"SPC_s = {fmt_num(spc_s,1)}; OAR_s = {fmt_num(oar_s,1)}; BDI_s = {fmt_num(background_design_score,1)}; ASL_s = {fmt_num(asl_s,1)}; raw average = ({fmt_num(spc_s,1)} + {fmt_num(oar_s,1)} + {fmt_num(background_design_score,1)} + {fmt_num(asl_s,1)}) / 4 = {fmt_num(fmi,1)}; guardrails = {', '.join(g['reason'] + ' => cap ' + fmt_num(g['cap'],1) for g in guardrails) if guardrails else 'none'}; final FMI = {fmt_num(final_fmi,1)}",
        fmi,
        final_fmi,
        'Interpretation: 100 means the feature is structurally maintainable; low values point to debt across steps, scenarios, Background design, or scenario length.'
    )

    report = {
        'source_name': aggregate.get('source_name', ''),
        'summary': {
            'feature_title': aggregate.get('feature_title'),
            'scenario_count': aggregate['scenario_count'],
            'outline_count': aggregate['outline_count'],
            'background_count': aggregate.get('background_count', 0),
            'background_step_count': aggregate['background_step_count'],
            'avg_background_step_count': aggregate.get('avg_background_step_count', 0.0),
            'total_step_count': aggregate['total_steps'],
            'unique_normalized_steps': aggregate['unique_normalized_steps'],
        },
        'metrics': metrics,
        'subscores': {
            'step_quality': round(spc_s, 1),
            'scenario_design': round((oar_s + asl_s) / 2.0, 1),
            'background_design': round(background_design_score, 1),
            'feature_maintainability_index': round(final_fmi, 1),
        },
        'guardrails_applied': guardrails,
        'weighted_total_before_caps': round(fmi, 1),
        'final_total': round(final_fmi, 1),
        'interpretation': interpretation_label(final_fmi),
        'total_row': total_row,
        'aggregate_counts': dict(aggregate),
        'report_kind': 'feature',
    }
    feature_key_source = report['summary'].get('feature_title') or report.get('source_name') or 'untitled-feature'
    report['state_feature_key'] = normalize_feature_key(feature_key_source)
    return report


def combine_aggregates(feature_reports: List[Dict]) -> Dict:
    aggregate = {
        'source_name': 'Totals',
        'feature_title': 'Suite totals',
        'scenario_count': 0,
        'outline_count': 0,
        'background_count': 0,
        'background_step_count': 0,
        'avg_background_step_count': 0.0,
        'max_background_step_count': 0,
        'total_steps': 0,
        'total_executed_steps': 0,
        'unique_normalized_steps': 0,
        'background_over6': False,
        'background_overused_features': 0,
        'total_scenario_steps': 0,
        'background_candidate_distinct_lines': 0,
        'background_candidate_occurrences': 0,
    }
    duplicate_occurrences = []
    for report in feature_reports:
        counts = report['aggregate_counts']
        aggregate['scenario_count'] += counts['scenario_count']
        aggregate['outline_count'] += counts['outline_count']
        aggregate['background_count'] += counts.get('background_count', 0)
        aggregate['background_step_count'] += counts['background_step_count']
        aggregate['max_background_step_count'] = max(aggregate['max_background_step_count'], counts.get('avg_background_step_count', counts['background_step_count']))
        aggregate['total_steps'] += counts.get('total_executed_steps', counts['total_steps'])
        aggregate['total_executed_steps'] += counts.get('total_executed_steps', counts['total_steps'])
        aggregate['background_over6'] = aggregate['background_over6'] or counts['background_over6']
        aggregate['total_scenario_steps'] += counts.get('total_scenario_steps', 0)
        aggregate['background_candidate_distinct_lines'] += counts.get('background_candidate_distinct_lines', 0)
        aggregate['background_candidate_occurrences'] += counts.get('background_candidate_occurrences', 0)
        aggregate['background_overused_features'] += counts.get('background_overused_features', 0)
        duplicate_occurrences.extend(counts['step_occurrences'])
    aggregate['unique_normalized_steps'] = len({occurrence['normalized'] for occurrence in duplicate_occurrences})
    aggregate['avg_background_step_count'] = safe_div(aggregate['background_step_count'], aggregate.get('background_count', 0), 0.0)
    aggregate['step_occurrences'] = duplicate_occurrences
    aggregate['suite_bdi'] = aggregate['background_candidate_distinct_lines'] * safe_div(aggregate['background_candidate_occurrences'], aggregate['total_scenario_steps'], 0.0)
    return aggregate


def report_title(report: Dict) -> str:
    return report['summary'].get('feature_title') or report.get('source_name') or 'Untitled feature'


def build_suite_report(aggregate: Dict, feature_reports: List[Dict]) -> Dict:
    total_features = len(feature_reports)
    feature_metric_values = {metric: [] for metric in [
        'Parameterization of Steps (SPC)',
        'Scenario Outline for Data-Driven Rules (OAR)',
        'Background Design Index (BDI)',
        'Average Scenario Length (ASL)',
    ]}
    feature_fmis = []
    for report in feature_reports:
        by_metric = {item['metric']: item for item in report['metrics']}
        for key in feature_metric_values:
            feature_metric_values[key].append(by_metric[key]['normalized_score'])
        feature_fmis.append(report['final_total'])

    def avg_feature_metric(name: str) -> float:
        return safe_div(sum(feature_metric_values[name]), total_features, 0.0)

    spc_s = avg_feature_metric('Parameterization of Steps (SPC)')
    oar_s = avg_feature_metric('Scenario Outline for Data-Driven Rules (OAR)')
    bdi_s = avg_feature_metric('Background Design Index (BDI)')
    asl_s = avg_feature_metric('Average Scenario Length (ASL)')

    def avg_calc(name: str, label: str) -> str:
        values = feature_metric_values[name]
        return f"Feature {label} values = {values}; {label}_suite = ({' + '.join(fmt_num(v,1) for v in values)}) / {total_features} = {fmt_num(safe_div(sum(values), total_features, 0.0),1)}" if total_features else f'No features available; {label}_suite defaults to 0'

    structural_index = safe_div(sum(feature_fmis), total_features, 0.0)
    suite_background_debt_raw = aggregate['suite_bdi']

    metrics = [
        metric_row('Parameterization of Steps (SPC)', 'suite', 'Average step parameterization quality across all analyzed features.', 'Take the normalized SPC_s score from each feature and compute the average.', 'SPC_suite = sum(feature SPC_s) / total features.', avg_calc('Parameterization of Steps (SPC)', 'SPC'), spc_s, spc_s, 'Interpretation: higher values indicate that the suite avoids duplicating step families with embedded literals.'),
        metric_row('Scenario Outline for Data-Driven Rules (OAR)', 'suite', 'Average quality of data-driven rule modeling with Scenario Outline across the suite.', 'Take the normalized OAR_s score from each feature and compute the average.', 'OAR_suite = sum(feature OAR_s) / total features.', avg_calc('Scenario Outline for Data-Driven Rules (OAR)', 'OAR'), oar_s, oar_s, 'Interpretation: higher values indicate that the suite consolidates repeated scenarios through Scenario Outline.'),
        metric_row('Background Design Index (BDI)', 'suite', 'Average Background design quality across all features, including overuse and repeated setup debt.', 'Take the normalized BDI_s score from each feature and compute the average. Also report raw Background diagnostics for the full suite.', 'BDI_suite = sum(feature BDI_s) / total features. Raw suite background debt = suite_background_candidate_distinct_lines * (suite_background_candidate_occurrences / suite_total_scenario_steps).', f"Feature BDI values = {feature_metric_values['Background Design Index (BDI)']}; average BDI_s = {fmt_num(bdi_s,1)}; overused Background features = {aggregate.get('background_overused_features', 0)}; suite background candidate distinct lines = {aggregate['background_candidate_distinct_lines']}; suite background candidate occurrences = {aggregate['background_candidate_occurrences']}; suite total scenario steps = {aggregate['total_scenario_steps']}; raw suite background debt = {fmt_num(suite_background_debt_raw)}", suite_background_debt_raw, bdi_s, 'Interpretation: lower values indicate heavy Background sections or repeated setup across several features.'),
        metric_row('Average Scenario Length (ASL)', 'suite', 'Average scenario length quality across all analyzed features.', 'Take the normalized ASL_s score from each feature and compute the average.', 'ASL_suite = sum(feature ASL_s) / total features.', avg_calc('Average Scenario Length (ASL)', 'ASL'), asl_s, asl_s, 'Interpretation: lower values indicate long scenarios that increase maintenance cost.'),
        metric_row('Structural Suite Index (SSI)', 'suite', 'Final structural maintainability index for the suite.', 'Take the final FMI of each feature and compute the average.', 'SSI = sum(feature FMI values) / total features', f"Feature FMI values = {[fmt_num(v,1) for v in feature_fmis]}; SSI = ({' + '.join(fmt_num(v,1) for v in feature_fmis)}) / {total_features} = {fmt_num(structural_index,1)}" if total_features else 'No features available; SSI defaults to 0', structural_index, structural_index, 'Interpretation: summarizes the average structural health of the suite.'),
    ]

    return {
        'source_name': 'Totals',
        'summary': {
            'feature_title': 'Suite totals',
            'scenario_count': aggregate['scenario_count'],
            'outline_count': aggregate['outline_count'],
            'background_count': aggregate.get('background_count', 0),
            'background_step_count': aggregate['background_step_count'],
            'avg_background_step_count': aggregate.get('avg_background_step_count', 0.0),
            'total_step_count': aggregate['total_steps'],
            'unique_normalized_steps': aggregate['unique_normalized_steps'],
        },
        'metrics': metrics,
        'subscores': {
            'structural_suite_index': round(structural_index, 1),
            'background_design': round(bdi_s, 1),
        },
        'guardrails_applied': [],
        'weighted_total_before_caps': None,
        'final_total': None,
        'interpretation': interpretation_label(structural_index),
        'aggregate_counts': dict(aggregate),
        'report_kind': 'suite',
    }


def build_duplicate_steps(feature_reports: List[Dict]) -> List[Dict]:
    occurrences = []
    for report in feature_reports:
        occurrences.extend(report['aggregate_counts']['step_occurrences'])
    by_norm: Dict[str, List[Dict]] = defaultdict(list)
    for occurrence in occurrences:
        by_norm[occurrence['normalized']].append(occurrence)

    rows = []
    for norm, occs in by_norm.items():
        if len(occs) <= 1:
            continue
        used_in_locations = sorted({f"{o['feature']} / {o['scenario']}" for o in occs})
        used_in = '; '.join(used_in_locations)
        amount = len(occs)
        for occurrence in occs:
            rows.append({
                'feature': occurrence['feature'],
                'scenario': occurrence['scenario'],
                'step': occurrence['step'],
                'used_in': used_in,
                'amount': amount,
                'normalized': norm,
            })
    rows.sort(key=lambda row: (row['feature'], row['scenario'], row['step']))
    return rows


def rows_for_report(report: Dict) -> List[Dict]:
    return list(report['metrics']) + ([report['total_row']] if report['report_kind'] == 'feature' else [])


def to_markdown_single(report: Dict, heading_level: int = 1) -> str:
    summary = report['summary']
    lines = [f"{'#' * heading_level} Gherkin Maintainability Report"]
    lines.append(f"\n**Feature:** {report_title(report)}")
    if report.get('source_name'):
        lines.append(f"**Source:** {report['source_name']}")
    lines.append(f"\n**Scenarios:** {summary['scenario_count']}  ")
    lines.append(f"**Scenario Outlines:** {summary['outline_count']}  ")
    lines.append(f"**Background Sections:** {summary.get('background_count', 0)}  ")
    lines.append(f"**Background Steps:** {summary['background_step_count']}  ")
    lines.append(f"**Average Background Steps:** {fmt_num(summary.get('avg_background_step_count', 0.0))}  ")
    lines.append(f"**Total Steps:** {summary['total_step_count']}  ")
    lines.append(f"**Unique Steps:** {summary['unique_normalized_steps']}\n")
    if report['report_kind'] == 'feature':
        lines.append('| Metric / Concept | Scope (step, scenario, feature, suite) | What it measures for maintainability | How to measure it | Practical formula / rule | Calculated Data | Raw value | Normalized score (0-100) | Example / interpretation | Raw Data |')
        lines.append('|---|---|---|---|---|---|---:|---:|---|---|')
        for metric in rows_for_report(report):
            lines.append(
                f"| {table_safe(metric['metric'])} | {table_safe(metric['concept_scope'])} | {table_safe(metric['what_it_measures'])} | {table_safe(metric['how_to_measure_it'])} | {table_safe(metric['practical_formula_rule'])} | {table_safe(metric['calculated_data'])} | {table_safe(metric['raw_value'])} | {table_safe(metric['normalized_score'])} | {table_safe(metric['example'])} | {table_safe(metric.get('raw_data', ''))} |"
            )
    else:
        lines.append('| Metric / Concept | Scope (step, scenario, feature, suite) | What it measures for maintainability | How to measure it | Practical formula / rule | Calculated Data | Raw value | Normalized score (0-100) | Example / interpretation |')
        lines.append('|---|---|---|---|---|---|---:|---:|---|')
        for metric in rows_for_report(report):
            lines.append(
                f"| {table_safe(metric['metric'])} | {table_safe(metric['concept_scope'])} | {table_safe(metric['what_it_measures'])} | {table_safe(metric['how_to_measure_it'])} | {table_safe(metric['practical_formula_rule'])} | {table_safe(metric['calculated_data'])} | {table_safe(metric['raw_value'])} | {table_safe(metric['normalized_score'])} | {table_safe(metric['example'])} |"
            )
    lines.append('\n## Summary scores')
    for key, value in report['subscores'].items():
        lines.append(f"- {key.replace('_', ' ').title()}: {value}")
    if report['report_kind'] == 'feature':
        lines.append(f"- Weighted total before caps: {report['weighted_total_before_caps']}")
        lines.append(f"- **Feature Maintainability Index: {report['final_total']} / 100 ({report['interpretation']})**")
    return '\n'.join(lines)


def to_markdown(feature_reports: List[Dict], suite_report: Dict = None, state_summary: Dict = None) -> str:
    if len(feature_reports) == 1 and suite_report is None:
        return to_markdown_single(feature_reports[0])
    blocks = ['# Gherkin Maintainability Suite Report']
    if state_summary:
        blocks.append(
            f"\nState summary: {state_summary['feature_count']} features in state; {state_summary['replaced_features']} replaced in this run; {state_summary['added_features']} added in this run."
        )
    for index, report in enumerate(feature_reports, start=1):
        blocks.append(f"\n## Feature {index}: {report_title(report)}")
        blocks.append(to_markdown_single(report, heading_level=3))
    if suite_report is not None:
        blocks.append('\n## Totals')
        blocks.append(to_markdown_single(suite_report, heading_level=3))
    return '\n\n'.join(blocks)


def unique_sheet_name(existing: set, label: str) -> str:
    invalid = '[]:*?/\\'
    clean = ''.join('_' if ch in invalid else ch for ch in (label or 'Feature'))
    clean = clean.strip() or 'Feature'
    clean = clean[:31]
    base = clean
    counter = 2
    while clean in existing:
        suffix = f"_{counter}"
        clean = (base[:31 - len(suffix)] + suffix)[:31]
        counter += 1
    existing.add(clean)
    return clean


def write_report_sheet(workbook, report: Dict, title: str) -> None:
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    sheet = workbook.create_sheet(title=title)
    title_font = Font(bold=True, size=14, color=LKS_NAVY)
    header_font = Font(bold=True, color=LKS_CHARCOAL)
    table_header_font = Font(bold=True, color=LKS_WHITE)
    total_font = Font(bold=True, color=LKS_CHARCOAL)
    thin = Side(style='thin', color=LKS_MID_GRAY)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill(fill_type='solid', fgColor=LKS_NAVY)
    section_fill = PatternFill(fill_type='solid', fgColor=LKS_LIGHT_GRAY)
    total_fill = PatternFill(fill_type='solid', fgColor=LKS_CYAN)
    wrap_top = Alignment(wrap_text=True, vertical='top')

    row = 1
    sheet.cell(row=row, column=1, value='Gherkin Maintainability Report').font = title_font
    row += 2
    summary = report['summary']
    summary_rows = [
        ('Feature', report_title(report)),
        ('Source', report.get('source_name') or ''),
        ('Scenarios', summary['scenario_count']),
        ('Scenario Outlines', summary['outline_count']),
        ('Background Sections', summary.get('background_count', 0)),
        ('Background Steps', summary['background_step_count']),
        ('Average Background Steps', fmt_num(summary.get('avg_background_step_count', 0.0))),
        ('Total Steps', summary['total_step_count']),
        ('Unique Steps', summary['unique_normalized_steps']),
    ]
    for label, value in summary_rows:
        sheet.cell(row=row, column=1, value=label).font = header_font
        sheet.cell(row=row, column=2, value=value)
        row += 1
    row += 1

    headers = ['Metric / Concept', 'Scope (step, scenario, feature, suite)', 'What it measures for maintainability', 'How to measure it', 'Practical formula / rule', 'Calculated Data', 'Raw value', 'Normalized score (0-100)', 'Example / interpretation', 'Raw Data']
    for col_idx, header in enumerate(headers, start=1):
        cell = sheet.cell(row=row, column=col_idx, value=header)
        cell.font = table_header_font
        cell.fill = header_fill
        cell.alignment = wrap_top
        cell.border = border
    row += 1

    for metric in rows_for_report(report):
        values = [metric['metric'], metric['concept_scope'], metric['what_it_measures'], metric['how_to_measure_it'], metric['practical_formula_rule'], metric['calculated_data'], metric['raw_value'], metric['normalized_score'], metric['example'], metric.get('raw_data', '')]
        for col_idx, value in enumerate(values, start=1):
            cell = sheet.cell(row=row, column=col_idx, value=value)
            cell.alignment = wrap_top
            cell.border = border
            if metric['metric'].endswith('(FMI)'):
                cell.fill = total_fill
                cell.font = total_font
        row += 1

    row += 1
    for label, value in report['subscores'].items():
        sheet.cell(row=row, column=1, value=label.replace('_', ' ').title()).font = header_font
        sheet.cell(row=row, column=2, value=value)
        row += 1
    sheet.cell(row=row, column=1, value='Weighted total before caps').font = header_font
    sheet.cell(row=row, column=2, value=report['weighted_total_before_caps'])
    row += 1
    sheet.cell(row=row, column=1, value='Feature Maintainability Index').font = header_font
    sheet.cell(row=row, column=2, value=report['final_total'])

    widths = {1: 34, 2: 18, 3: 42, 4: 38, 5: 42, 6: 64, 7: 12, 8: 22, 9: 32, 10: 56}
    for col_idx, width in widths.items():
        sheet.column_dimensions[get_column_letter(col_idx)].width = width
    for excel_row in sheet.iter_rows(min_row=1, max_row=sheet.max_row, min_col=1, max_col=10):
        for cell in excel_row:
            cell.alignment = wrap_top


def write_totals_sheet(workbook, feature_reports: List[Dict], suite_report: Dict) -> None:
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    sheet = workbook.create_sheet(title='Totals')
    title_font = Font(bold=True, size=14, color=LKS_NAVY)
    header_font = Font(bold=True, color=LKS_CHARCOAL)
    table_header_font = Font(bold=True, color=LKS_WHITE)
    total_font = Font(bold=True, color=LKS_CHARCOAL)
    thin = Side(style='thin', color=LKS_MID_GRAY)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill(fill_type='solid', fgColor=LKS_NAVY)
    section_fill = PatternFill(fill_type='solid', fgColor=LKS_LIGHT_GRAY)
    ssi_fill = PatternFill(fill_type='solid', fgColor=LKS_CYAN)
    wrap_top = Alignment(wrap_text=True, vertical='top')

    row = 1
    sheet.cell(row=row, column=1, value='Gherkin Maintainability Totals').font = title_font
    row += 2
    sheet.cell(row=row, column=1, value='Summary by feature').font = header_font
    row += 1
    summary_headers = ['Feature / Sheet', 'SPC_s', 'OAR_s', 'BDI_s', 'ASL_s', 'Feature MI']
    for col_idx, header in enumerate(summary_headers, start=1):
        cell = sheet.cell(row=row, column=col_idx, value=header)
        cell.font = table_header_font
        cell.fill = header_fill
        cell.border = border
        cell.alignment = wrap_top
    row += 1
    for report in feature_reports:
        by_metric = {m['metric']: m for m in report['metrics']}
        values = [
            report_title(report),
            by_metric['Parameterization of Steps (SPC)']['normalized_score'],
            by_metric['Scenario Outline for Data-Driven Rules (OAR)']['normalized_score'],
            by_metric['Background Design Index (BDI)']['normalized_score'],
            by_metric['Average Scenario Length (ASL)']['normalized_score'],
            report['final_total'],
        ]
        for col_idx, value in enumerate(values, start=1):
            cell = sheet.cell(row=row, column=col_idx, value=value)
            cell.border = border
            cell.alignment = wrap_top
        row += 1

    row += 2
    sheet.cell(row=row, column=1, value='Suite totals').font = header_font
    row += 1
    headers = ['Metric / Concept', 'Scope (step, scenario, feature, suite)', 'What it measures for maintainability', 'How to measure it', 'Practical formula / rule', 'Calculated Data', 'Normalized score (0-100)', 'Example / interpretation']
    for col_idx, header in enumerate(headers, start=1):
        cell = sheet.cell(row=row, column=col_idx, value=header)
        cell.font = table_header_font
        cell.fill = header_fill
        cell.alignment = wrap_top
        cell.border = border
    row += 1
    for metric in suite_report['metrics']:
        values = [metric['metric'], metric['concept_scope'], metric['what_it_measures'], metric['how_to_measure_it'], metric['practical_formula_rule'], metric['calculated_data'], metric['normalized_score'], metric['example']]
        for col_idx, value in enumerate(values, start=1):
            cell = sheet.cell(row=row, column=col_idx, value=value)
            cell.border = border
            cell.alignment = wrap_top
            if metric['metric'] == 'Structural Suite Index (SSI)':
                cell.fill = ssi_fill
                cell.font = total_font
        row += 1

    widths = {1: 36, 2: 20, 3: 42, 4: 38, 5: 46, 6: 72, 7: 16, 8: 34}
    for col_idx, width in widths.items():
        sheet.column_dimensions[get_column_letter(col_idx)].width = width
    for excel_row in sheet.iter_rows(min_row=1, max_row=sheet.max_row, min_col=1, max_col=8):
        for cell in excel_row:
            cell.alignment = wrap_top


def write_duplicate_steps_sheet(workbook, duplicate_rows: List[Dict]) -> None:
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    sheet = workbook.create_sheet(title='Repeated Steps')
    title_font = Font(bold=True, size=14, color=LKS_NAVY)
    header_font = Font(bold=True, color=LKS_WHITE)
    thin = Side(style='thin', color=LKS_MID_GRAY)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill(fill_type='solid', fgColor=LKS_NAVY)
    wrap_top = Alignment(wrap_text=True, vertical='top')

    row = 1
    sheet.cell(row=row, column=1, value='Repeated / Non-Unique Steps Across Suite').font = title_font
    row += 2
    headers = ['feature', 'scenario', 'step', 'used in', 'amount']
    for col_idx, header in enumerate(headers, start=1):
        cell = sheet.cell(row=row, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = border
        cell.alignment = wrap_top
    row += 1
    for item in duplicate_rows:
        values = [item['feature'], item['scenario'], item['step'], item['used_in'], item['amount']]
        for col_idx, value in enumerate(values, start=1):
            cell = sheet.cell(row=row, column=col_idx, value=value)
            cell.border = border
            cell.alignment = wrap_top
        row += 1

    widths = {1: 28, 2: 32, 3: 48, 4: 80, 5: 10}
    for col_idx, width in widths.items():
        sheet.column_dimensions[get_column_letter(col_idx)].width = width


def clean_zip_feature_title(name: str) -> str:
    stem = Path(name).stem
    stem = re.sub(r"\s*\(\d+\)$", "", stem)
    stem = stem.replace('_', ' ').replace('-', ' ')
    stem = re.sub(r"\s+", " ", stem).strip()
    return stem.title() or stem or 'Zip Feature'


def aggregate_parsed_features(features: List[Dict], title: str) -> Dict:
    aggregate = {"title": title, "background_steps": [], "background_count": 0, "scenarios": []}
    for feature in features:
        aggregate['background_steps'].extend(feature.get('background_steps', []))
        aggregate['background_count'] += feature.get('background_count', 0)
        aggregate['scenarios'].extend(feature.get('scenarios', []))
    return aggregate


def load_zip_as_single_feature(path: Path) -> Dict:
    parsed_features: List[Dict] = []
    with zipfile.ZipFile(path) as zf:
        for info in sorted(zf.infolist(), key=lambda item: item.filename.lower()):
            if info.is_dir():
                continue
            member = Path(info.filename)
            if member.suffix.lower() not in VALID_SUFFIXES:
                continue
            raw = zf.read(info.filename)
            text = raw.decode('utf-8', errors='replace')
            parsed_features.append(parse_gherkin(text))
    if not parsed_features:
        raise SystemExit(f'No supported Gherkin files found inside zip: {path}')
    return {'source_name': path.name, 'parsed_feature': aggregate_parsed_features(parsed_features, clean_zip_feature_title(path.name))}


def export_xlsx(feature_reports: List[Dict], suite_report: Dict, duplicate_rows: List[Dict], output_path: str) -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.remove(workbook.active)
    existing = set()
    for report in feature_reports:
        write_report_sheet(workbook, report, unique_sheet_name(existing, report_title(report)))
    write_totals_sheet(workbook, feature_reports, suite_report)
    write_duplicate_steps_sheet(workbook, duplicate_rows)
    workbook.save(output_path)


def resolve_inputs(inputs: List[str]) -> List[Dict]:
    resolved = []
    raw_texts = []
    for item in inputs:
        path = Path(item)
        if path.exists():
            if path.is_dir():
                for child in sorted(path.rglob('*')):
                    if child.is_file() and child.suffix.lower() in VALID_SUFFIXES:
                        resolved.append({'source_name': child.name, 'text': child.read_text(encoding='utf-8')})
            elif path.is_file() and path.suffix.lower() == '.zip':
                resolved.append(load_zip_as_single_feature(path))
            elif path.is_file() and path.suffix.lower() in VALID_SUFFIXES:
                resolved.append({'source_name': path.name, 'text': path.read_text(encoding='utf-8')})
        else:
            raw_texts.append(item)
    if raw_texts:
        resolved.append({'source_name': 'raw_input', 'text': '\n'.join(raw_texts)})
    return resolved


def load_state(state_path: str) -> Dict:
    if not state_path:
        return {'version': STATE_VERSION, 'features': {}, 'meta': {}}
    path = Path(state_path)
    if not path.exists():
        raise SystemExit(f'State file not found: {state_path}')
    data = json.loads(path.read_text(encoding='utf-8'))
    if not isinstance(data, dict):
        raise SystemExit('State file must contain a JSON object.')
    if data.get('version') != STATE_VERSION:
        raise SystemExit(f"Unsupported state version: {data.get('version')}")
    features = data.get('features', {})
    if not isinstance(features, dict):
        raise SystemExit('State file field "features" must be an object.')
    return data


def save_state(state: Dict, output_path: str) -> None:
    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(state, indent=2, ensure_ascii=False), encoding='utf-8')


def sort_reports(feature_reports: List[Dict]) -> List[Dict]:
    return sorted(feature_reports, key=lambda report: (report_title(report).lower(), report.get('source_name', '').lower()))


def state_feature_reports(state: Dict) -> List[Dict]:
    features = state.get('features', {})
    reports = []
    for key, payload in features.items():
        report = payload.get('report', payload)
        clean_report = json.loads(json.dumps(report))
        clean_report['state_feature_key'] = key
        reports.append(clean_report)
    return sort_reports(reports)


def main() -> None:
    parser = argparse.ArgumentParser(description='Analyze Gherkin maintainability metrics with optional incremental JSON state.')
    parser.add_argument('inputs', nargs='*', help='One or more Gherkin files, directories, or raw Gherkin text snippets.')
    parser.add_argument('--format', choices=['json', 'markdown'], default='json')
    parser.add_argument('--excel-output', help='Optional path to write an Excel (.xlsx) report with one sheet per feature plus Totals and Repeated Steps.')
    parser.add_argument('--state-in', help='Optional JSON state file from a previous batch.')
    parser.add_argument('--state-out', help='Optional path to write the updated JSON state for the next batch.')
    args = parser.parse_args()

    input_items = resolve_inputs(args.inputs)
    if not input_items and not args.state_in:
        raise SystemExit('Provide at least one Gherkin input or a --state-in file.')

    loaded_state = load_state(args.state_in)
    features_state = loaded_state.get('features', {})
    if not isinstance(features_state, dict):
        raise SystemExit('State file field "features" must be an object.')

    new_feature_reports = []
    replaced_features = 0
    added_features = 0
    for item in input_items:
        feature = item.get('parsed_feature') or parse_gherkin(item['text'])
        aggregate = summarize_feature(feature, item['source_name'])
        report = build_feature_report(aggregate)
        feature_key = report['state_feature_key']
        if feature_key in features_state:
            replaced_features += 1
        else:
            added_features += 1
        features_state[feature_key] = {
            'feature_key': feature_key,
            'feature_title': report['summary'].get('feature_title'),
            'source_name': report.get('source_name', ''),
            'report': strip_report_for_state(report),
        }
        new_feature_reports.append(report)

    all_feature_reports = state_feature_reports({'features': features_state})
    suite_aggregate = combine_aggregates(all_feature_reports)
    suite_report = build_suite_report(suite_aggregate, all_feature_reports)
    duplicate_rows = build_duplicate_steps(all_feature_reports)

    state = {
        'version': STATE_VERSION,
        'features': features_state,
        'meta': {
            'feature_count': len(features_state),
            'last_run_utc': datetime.now(timezone.utc).isoformat(),
            'added_features': added_features,
            'replaced_features': replaced_features,
        },
    }

    if args.state_out:
        save_state(state, args.state_out)

    if args.excel_output:
        export_xlsx(all_feature_reports, suite_report, duplicate_rows, args.excel_output)

    state_summary = {
        'feature_count': len(features_state),
        'added_features': added_features,
        'replaced_features': replaced_features,
    }
    payload = {
        'reports': all_feature_reports,
        'new_reports': sort_reports(new_feature_reports),
        'suite_report': suite_report,
        'duplicate_steps': duplicate_rows,
        'state_summary': state_summary,
    }
    if args.state_out:
        payload['state_output'] = args.state_out
    if args.excel_output:
        payload['excel_output'] = args.excel_output

    if args.format == 'markdown':
        print(to_markdown(all_feature_reports, suite_report, state_summary=state_summary))
    else:
        json.dump(payload, sys.stdout, indent=2, ensure_ascii=False)
        print()


if __name__ == '__main__':
    main()
